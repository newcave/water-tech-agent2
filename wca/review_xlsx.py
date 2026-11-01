# -*- coding: utf-8 -*-
"""기술 지도 검토용 엑셀 — 첫 해 → 고른 해 사분면 변화를 연구소별 표로 내려받아 사람이 '판단'을 채워 돌려주게 한다.
(화면의 가중치·연도 그대로. 비공개 문서를 공유하기 어려울 때 쓰는 내려받기 — 2026-11-01)"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

INST_NAME = {"INST-01": "경영연구소", "INST-02": "수자원환경연구소", "INST-03": "상하수도연구소", "INST-04": "물인프라안전연구소",
             "INST-05": "물에너지연구소", "INST-06": "수자원위성연구소", "INST-07": "AI연구소"}
QN = {"core": "주력", "gap": "실용화 과제", "applied": "실용화 성숙", "weak": "위크시그널", "review": "검토 필요"}
QORDER = ["core", "gap", "applied", "weak", "review"]
RAW = {"V": ("논문 수(3년)", "n3", "{:,.0f}편"), "G": ("연 성장", "g", "{:+.0f}%"), "Q": ("상위10% 비율", "q", "{:.0f}%"),
       "J": ("저널 IF", "j", "{:.1f}"), "I": ("산업계 공저", "i", "{:.1f}%")}

HEAD = PatternFill("solid", fgColor="1F3A5F")
SUB = PatternFill("solid", fgColor="EEF3F8")
MOVED = PatternFill("solid", fgColor="FFF7E6")
INPUT = PatternFill("solid", fgColor="FFFDE7")
THIN = Side(style="thin", color="D0D7E2")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fmt(f, v):
    if v is None:
        return "–"
    s = f.format(v)
    return s.replace("-0%", "0%").replace("+0%", "0%")


def _drivers(u, a, b, k=2):
    c0, c1 = (u.get("c") or {}).get(a, {}), (u.get("c") or {}).get(b, {})
    r0, r1 = (u.get("r") or {}).get(a, {}), (u.get("r") or {}).get(b, {})
    d = sorted(((abs((c1.get(x) or 0) - (c0.get(x) or 0)), x) for x in RAW
                if c0.get(x) is not None and c1.get(x) is not None), reverse=True)
    out = []
    for _, x in d[:k]:
        nm, key, f = RAW[x]
        if r0.get(key) is not None and r1.get(key) is not None:
            out.append(f"{nm} {_fmt(f, r0[key])}→{_fmt(f, r1[key])}")
    return " · ".join(out)


def _header(ws, row, cols):
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def build(sc: dict, S: dict, a: str, b: str, y_basis: str) -> bytes:
    """sc = scores.json, S = {unit: {year: (x, y, quad)}} (화면 가중치로 다시 매긴 값), a → b 연도."""
    units = sc.get("units") or {}
    ids = sorted(units, key=lambda u: (units[u]["inst"], u))
    q = {u: (S[u][a][2], S[u][b][2]) for u in ids}
    wb = Workbook()

    # ── 1. 읽는 법 ──
    ws = wb.active
    ws.title = "읽는 법"
    moved = sum(1 for u in ids if q[u][0] != q[u][1])
    lines = [
        (f"기술 지도 추세 변화 검토표 ({a} → {b})", Font(bold=True, size=14)),
        (f"41개 세부 기술 중 {moved}개가 사분면을 옮겼고 {len(ids) - moved}개는 그대로입니다.", Font(size=11)),
        ("", None),
        ("검토 방법", Font(bold=True)),
        ("· '연구소별 검토' 시트의 노란 칸에 적어 주세요: 판단(맞다·다르다·모르겠다), 검토 의견, 검토자.", None),
        ("· 판단은 '이 움직임이 현장 감각과 맞는가'입니다. 다르면 의견에 한 줄이면 충분합니다.", None),
        ("", None),
        ("읽는 법", Font(bold=True)),
        ("· 점수는 그해 41개 세부 기술 사이의 순위(0~100)입니다. 50이 가운데라 절대 수준이 아니라 서로 견준 위치입니다.", None),
        ("· 논문 점수 = 논문 수(3년)·연 성장(5년)·피인용 상위 10% 비율(3년)·게재 저널 IF(2년). 자료는 OpenAlex 학술 논문.", None),
        ("· 실용화 점수 = " + ("특허(WCS) 규모·성장·등록률·기업 출원 + 산업계 공저." if y_basis == "patents" else
                           "아직 특허가 아니라 기업 소속 저자가 참여한 논문 비율 하나로 매긴 임시 값입니다."), None),
        ("· 사분면: 주력(둘 다 위) · 실용화 과제(논문 위, 실용화 아래) · 실용화 성숙(논문 아래, 실용화 위) · "
         "검토 필요(둘 다 아래) · 위크시그널(둘 다 아래지만 성장 상위 30%).", None),
        ("· '표본 작음' = 어느 한 해라도 3년 논문이 50편 미만. 몇 편 차이로 순위가 크게 흔들립니다.", None),
        ("· 2025년은 OpenAlex가 자료 저장소·작은 저널을 많이 더 담아 정책·회복력 같은 주제의 논문 수와 성장이 부풀었을 수 있습니다.", None),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        if f:
            c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120

    # ── 2. 한눈에 ──
    ws = wb.create_sheet("한눈에")
    _header(ws, 1, ["사분면", a, b, "변화"])
    for i, k in enumerate(QORDER, start=2):
        n0 = sum(1 for u in ids if q[u][0] == k)
        n1 = sum(1 for u in ids if q[u][1] == k)
        for j, v in enumerate([QN[k], n0, n1, n1 - n0], start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BOX
            c.alignment = Alignment(horizontal="center")
    r0 = len(QORDER) + 3
    ws.cell(row=r0, column=1, value=f"{a} 위치(줄) → {b} 위치(칸)").font = Font(bold=True)
    _header(ws, r0 + 1, [f"{a} \\ {b}"] + [QN[k] for k in QORDER])
    for i, ka in enumerate(QORDER):
        ws.cell(row=r0 + 2 + i, column=1, value=QN[ka]).font = Font(bold=True)
        ws.cell(row=r0 + 2 + i, column=1).border = BOX
        for j, kb in enumerate(QORDER, start=2):
            n = sum(1 for u in ids if q[u] == (ka, kb))
            c = ws.cell(row=r0 + 2 + i, column=j, value=n)
            c.border = BOX
            c.alignment = Alignment(horizontal="center")
            if ka == kb:
                c.font = Font(bold=True)
                c.fill = SUB
    for col, w in zip("ABCDEF", [16, 12, 12, 12, 12, 12]):
        ws.column_dimensions[col].width = w

    # ── 3. 연구소별 검토 ──
    ws = wb.create_sheet("연구소별 검토")
    cols = ["연구소", "세부 기술", "검색어(OpenAlex)", "표본 작음", f"{a} 사분면", f"{b} 사분면", "이동",
            f"논문 점수 {a}", f"논문 점수 {b}", f"실용화 점수 {a}", f"실용화 점수 {b}", "크게 바뀐 지표", "판단", "검토 의견", "검토자"]
    _header(ws, 1, cols)
    dv = DataValidation(type="list", formula1='"맞다,다르다,모르겠다"', allow_blank=True, showDropDown=False)
    dv.prompt, dv.promptTitle = "맞다 · 다르다 · 모르겠다", "판단"
    ws.add_data_validation(dv)
    for i, uid in enumerate(ids, start=2):
        u = units[uid]
        (x0, y0, _), (x1, y1, _) = S[uid][a], S[uid][b]
        n_small = min(((u.get("r") or {}).get(a, {}).get("n3") or 0), ((u.get("r") or {}).get(b, {}).get("n3") or 0)) < 50
        qa, qb = q[uid]
        row = [INST_NAME.get(u["inst"], u["inst"]), u["name"], u.get("en", ""), "예" if n_small else "",
               QN.get(qa, "–"), QN.get(qb, "–"), "그대로" if qa == qb else f"{QN.get(qa, '–')} → {QN.get(qb, '–')}",
               None if x0 is None else round(x0), None if x1 is None else round(x1),
               None if y0 is None else round(y0), None if y1 is None else round(y1),
               _drivers(u, a, b), None, None, None]
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BOX
            c.alignment = Alignment(vertical="center", wrap_text=j in (12, 14))
            if j >= 13:
                c.fill = INPUT
            elif qa != qb:
                c.fill = MOVED
        dv.add(f"M{i}")
    widths = [16, 20, 30, 9, 12, 12, 24, 10, 10, 11, 11, 44, 11, 40, 10]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(ids) + 1}"
    ws.row_dimensions[1].height = 32

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
